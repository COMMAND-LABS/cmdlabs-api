"""
Upload router for Vector Stores module.
Handles file uploads to cloud storage and triggers async processing via Pub/Sub.
"""
import io
import json
import logging
from fastapi import APIRouter, Request, UploadFile, File, HTTPException, Form
from typing import Any, Dict, Optional
from starlette.datastructures import Headers
from src.deps import org_dependency, jwt_dependency, db_dependency, ensure_account, account_id_from_claims
from src.services.vector_stores_upload_service import (
    QNA_INGEST_TOPIC,
    TXT_INGEST_TOPIC,
    VectorStoresUploadService,
)
from src.services.vector_store_access import authorize_vector_store
from src.services.account_gcs_service import AccountGcsCredentialMissing
from src.services.ingestion_log import record_ingestion_log
import uuid
from src.rate_limit import limiter

logger = logging.getLogger(__name__)

router = APIRouter()


def _authorize_upload(db, decoded_jwt, org, index_name: str, owner_account_id: Optional[int]) -> int:
    """Authenticate the caller and resolve the knowledge base owner to write as."""
    if not decoded_jwt:
        raise HTTPException(status_code=401, detail="Authentication required")

    caller_account_id = account_id_from_claims(decoded_jwt)
    # Ingesting into a shared knowledge base requires write (admin) access; for
    # your own KB this returns you unchanged. All GCS/Pinecone/log writes use the owner.
    account_id = authorize_vector_store(db, caller_account_id, index_name, owner_account_id, require_write=True, org_id=org.org_id)

    # Validate account exists
    ensure_account(db, account_id)
    return account_id


def _require_index_and_namespace(index_name: str, namespace: str) -> None:
    """Validate index_name is provided and namespace is present (it may be empty)."""
    if not index_name or not index_name.strip():
        raise HTTPException(
            status_code=400,
            detail="index_name is required"
        )

    # An empty namespace is Pinecone's default section (shown as "Main"), so it
    # is a valid target. Only a missing value is rejected.
    if namespace is None:
        raise HTTPException(
            status_code=400,
            detail="namespace is required"
        )


async def _upload_and_log(
    *,
    file: UploadFile,
    account_id: int,
    decoded_jwt: dict,
    index_name: str,
    namespace: str,
    comment: Optional[str],
    batch_number: Optional[str],
    db,
    request: Optional[Request],
    topic_name: str,
    log_prefix: str,
    extra_message_fields: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Store the file in the account's GCS bucket, publish it to ``topic_name``
    and record a PENDING ingestion log entry. Shared by every upload route."""
    # Generate batch number if not provided
    if batch_number is None:
        batch_number = str(uuid.uuid4())

    # Initialize upload service
    upload_service = VectorStoresUploadService()

    # Upload file to the account's GCS bucket and publish to Pub/Sub
    try:
        result = await upload_service.upload_file_and_publish(
            file=file,
            user_id=str(account_id),
            user_email=str(decoded_jwt.get('email', '')),
            index_name=index_name.strip(),
            namespace=namespace.strip(),
            jwt=request.cookies.get("jwt") if request else None,
            db=db,
            account_id=account_id,
            topic_name=topic_name,
            batch_number=batch_number,
            comment=comment,
            extra_message_fields=extra_message_fields,
        )
    except AccountGcsCredentialMissing as e:
        raise HTTPException(status_code=400, detail=str(e))

    if not result.get("success"):
        raise HTTPException(
            status_code=500,
            detail=result.get("error", "Failed to upload file")
        )

    # Create ingestion log entry with PENDING status. Best-effort: don't fail
    # the upload if logging fails.
    log_id = record_ingestion_log(
        db,
        log_prefix=log_prefix,
        refresh=True,
        account_id=account_id,
        provider="pinecone",  # Default to pinecone, could be made configurable
        index_name=index_name.strip(),
        namespace=namespace.strip(),
        filenames=[file.filename],
        comment=comment,
        gcs_bucket=result.get("gcs_bucket"),
        gcs_file_path=result.get("gcs_file_path"),
        operation_type='INGEST',  # Direct string assignment for PostgreSQL enum
        status='PENDING',  # Direct string assignment for PostgreSQL enum
        vectors_added=0,
        vectors_deleted=0,
        vectors_failed=0,
        batch_number=batch_number
    )
    if log_id is not None:
        result["log_id"] = log_id

    return result


@router.post("/upload-csv")
@limiter.limit("100/minute")
async def upload_csv_file(
    file: UploadFile = File(..., description="CSV file to upload"),
    index_name: str = Form(..., description="Pinecone index name"),
    namespace: str = Form("", description="Pinecone namespace (empty for the default \"Main\" section)"),
    comment: Optional[str] = Form(None, description="Optional comment for the ingestion log"),
    batch_number: Optional[str] = Form(None, description="Optional batch UUID for grouping related operations"),
    owner_account_id: Optional[int] = Form(None, description="Owner of a shared knowledge base to ingest into (requires write/admin access)"),
    db: db_dependency = None,
    decoded_jwt: jwt_dependency = None,
    org: org_dependency = None,
    request: Request = None
):
    """
    Upload a CSV file to Google Cloud Storage and queue it for async processing.
    The file will be processed (chunked and uploaded to Pinecone) asynchronously via cloud function.

    Creates an ingestion log entry with PENDING status that will be updated when processing completes.
    """
    account_id = _authorize_upload(db, decoded_jwt, org, index_name, owner_account_id)

    # Validate file type
    if not file.filename or not file.filename.endswith('.csv'):
        raise HTTPException(
            status_code=400,
            detail="Only .csv files are supported"
        )

    _require_index_and_namespace(index_name, namespace)

    return await _upload_and_log(
        file=file,
        account_id=account_id,
        decoded_jwt=decoded_jwt,
        index_name=index_name,
        namespace=namespace,
        comment=comment,
        batch_number=batch_number,
        db=db,
        request=request,
        topic_name=QNA_INGEST_TOPIC,
        log_prefix="[UPLOAD CSV]",
    )

@router.post("/upload-pdf-faq")
@limiter.limit("100/minute")
async def upload_pdf_faq(
    file: UploadFile = File(..., description="Original PDF to store as the source document"),
    index_name: str = Form(..., description="Pinecone index name"),
    namespace: str = Form("", description="Pinecone namespace (empty for the default \"Main\" section)"),
    qna_pairs: str = Form(..., description="JSON array of reviewed Q&A pairs: [{\"q\": str, \"a\": str}]"),
    comment: Optional[str] = Form(None, description="Optional comment for the ingestion log"),
    batch_number: Optional[str] = Form(None, description="Optional batch UUID for grouping related operations"),
    owner_account_id: Optional[int] = Form(None, description="Owner of a shared knowledge base to ingest into (requires write/admin access)"),
    db: db_dependency = None,
    decoded_jwt: jwt_dependency = None,
    org: org_dependency = None,
    request: Request = None
):
    """
    Store the ORIGINAL PDF in Google Cloud Storage and queue the reviewed Q&A
    pairs for async ingestion.

    Unlike upload-csv, the file stored in GCS is the source PDF (not the data
    being parsed). The Q&A pairs ride in the Pub/Sub message; the ingest cloud
    function builds one vector per pair and stamps each vector's source metadata
    (storage_bucket / storage_path / filename) to the PDF, so every FAQ entry can
    be traced back to its original source.
    """
    account_id = _authorize_upload(db, decoded_jwt, org, index_name, owner_account_id)

    # Validate file type
    if not file.filename or not file.filename.lower().endswith('.pdf'):
        raise HTTPException(status_code=400, detail="Only .pdf files are supported")

    _require_index_and_namespace(index_name, namespace)

    # Validate and normalize the Q&A pairs
    try:
        parsed_pairs = json.loads(qna_pairs)
    except (json.JSONDecodeError, TypeError) as e:
        raise HTTPException(status_code=400, detail=f"qna_pairs must be valid JSON: {e}")

    if not isinstance(parsed_pairs, list) or len(parsed_pairs) == 0:
        raise HTTPException(status_code=400, detail="qna_pairs must be a non-empty array")

    normalized_pairs = []
    for pair in parsed_pairs:
        if not isinstance(pair, dict):
            continue
        q = str(pair.get("q", "")).strip()
        a = str(pair.get("a", "")).strip()
        if q and a:
            normalized_pairs.append({"q": q, "a": a})

    if not normalized_pairs:
        raise HTTPException(status_code=400, detail="qna_pairs contained no valid {q, a} entries")

    # Store the PDF in the account's GCS bucket and publish to Pub/Sub. The
    # reviewed Q&A pairs ride along in the message via extra_message_fields.
    return await _upload_and_log(
        file=file,
        account_id=account_id,
        decoded_jwt=decoded_jwt,
        index_name=index_name,
        namespace=namespace,
        comment=comment,
        batch_number=batch_number,
        db=db,
        request=request,
        topic_name=QNA_INGEST_TOPIC,
        log_prefix="[UPLOAD PDF FAQ]",
        extra_message_fields={
            "source_type": "pdf_faq",
            "qna_pairs": normalized_pairs,
        },
    )


@router.post("/upload-text")
@limiter.limit("100/minute")
async def upload_text_file(
    file: UploadFile = File(..., description="Text file to upload (.txt or .md)"),
    index_name: str = Form(..., description="Pinecone index name"),
    namespace: str = Form("", description="Pinecone namespace (empty for the default \"Main\" section)"),
    comment: Optional[str] = Form(None, description="Optional comment for the ingestion log"),
    batch_number: Optional[str] = Form(None, description="Optional batch UUID for grouping related operations"),
    owner_account_id: Optional[int] = Form(None, description="Owner of a shared knowledge base to ingest into (requires write/admin access)"),
    db: db_dependency = None,
    decoded_jwt: jwt_dependency = None,
    org: org_dependency = None,
    request: Request = None
):
    """
    Upload a text file (.txt or .md) to Google Cloud Storage and queue it for async processing.
    The file will be processed (chunked and uploaded to Pinecone) asynchronously via cloud function.

    Creates an ingestion log entry with PENDING status that will be updated when processing completes.
    """
    account_id = _authorize_upload(db, decoded_jwt, org, index_name, owner_account_id)

    # Validate file type
    if not file.filename or not file.filename.endswith(('.txt', '.md')):
        raise HTTPException(
            status_code=400,
            detail="Only .txt and .md files are supported"
        )

    _require_index_and_namespace(index_name, namespace)

    return await _upload_and_log(
        file=file,
        account_id=account_id,
        decoded_jwt=decoded_jwt,
        index_name=index_name,
        namespace=namespace,
        comment=comment,
        batch_number=batch_number,
        db=db,
        request=request,
        # Free text goes to the TXT ingest function, which chunks the
        # document and reads YAML front matter. The Q&A function would
        # not recognise this file at all.
        topic_name=TXT_INGEST_TOPIC,
        log_prefix="[UPLOAD TEXT]",
    )


# Documents are read in full to extract their text, so cap them the same way the
# general file upload does.
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
DOCUMENT_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")
NO_TEXT_FOUND = "We couldn't find any text in this file — is it a scanned image?"


def _extract_pdf_text(data: bytes) -> str:
    """Plain text of every page, in order. Pages with no text are skipped."""
    import fitz  # PyMuPDF

    parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            text = page.get_text("text")
            if text and text.strip():
                parts.append(text.strip())
    return "\n\n".join(parts)


def _extract_docx_text(data: bytes) -> str:
    """Paragraph text followed by table cell text (one row per line)."""
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                text = cell.text.strip()
                # Merged cells repeat across the row; keep one copy.
                if text and (not cells or cells[-1] != text):
                    cells.append(text)
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


@router.post("/upload-document")
@limiter.limit("100/minute")
async def upload_document_file(
    file: UploadFile = File(..., description="Document to upload (.pdf, .docx, .txt or .md)"),
    index_name: str = Form(..., description="Pinecone index name"),
    namespace: str = Form("", description="Pinecone namespace (empty for the default \"Main\" section)"),
    comment: Optional[str] = Form(None, description="Optional comment for the ingestion log"),
    batch_number: Optional[str] = Form(None, description="Optional batch UUID for grouping related operations"),
    owner_account_id: Optional[int] = Form(None, description="Owner of a shared knowledge base to ingest into (requires write/admin access)"),
    db: db_dependency = None,
    decoded_jwt: jwt_dependency = None,
    org: org_dependency = None,
    request: Request = None
):
    """
    Upload a document for free-text ingestion.

    .txt and .md files are stored as-is (same as /upload-text). For .pdf and
    .docx the text is extracted here and stored as "<name>.txt", so the TXT
    ingest function chunks it exactly like an uploaded text file. The original
    filename is recorded in the ingestion log comment.
    """
    account_id = _authorize_upload(db, decoded_jwt, org, index_name, owner_account_id)

    filename = file.filename or ""
    lower = filename.lower()
    if not lower.endswith(DOCUMENT_EXTENSIONS):
        raise HTTPException(
            status_code=400,
            detail="Only PDF (.pdf), Word (.docx), and text (.txt, .md) files are supported",
        )

    _require_index_and_namespace(index_name, namespace)

    data = await file.read()
    if len(data) > MAX_DOCUMENT_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File exceeds the {MAX_DOCUMENT_BYTES // (1024 * 1024)} MB limit",
        )

    if lower.endswith((".txt", ".md")):
        if not data.strip():
            raise HTTPException(status_code=400, detail=NO_TEXT_FOUND)
        await file.seek(0)
        upload = file
        upload_comment = comment
    else:
        extractor = _extract_pdf_text if lower.endswith(".pdf") else _extract_docx_text
        try:
            text = extractor(data)
        except Exception:
            logger.exception("[UPLOAD DOCUMENT] Could not read %s", filename)
            raise HTTPException(
                status_code=400,
                detail="We couldn't read this file. It may be damaged or password-protected.",
            )
        if not text.strip():
            raise HTTPException(status_code=400, detail=NO_TEXT_FOUND)

        stem = filename.rsplit(".", 1)[0] or "document"
        upload = UploadFile(
            file=io.BytesIO(text.encode("utf-8")),
            filename=f"{stem}.txt",
            headers=Headers({"content-type": "text/plain; charset=utf-8"}),
        )
        note = f"Text extracted from {filename}"
        upload_comment = f"{comment.strip()} ({note})" if comment and comment.strip() else note

    return await _upload_and_log(
        file=upload,
        account_id=account_id,
        decoded_jwt=decoded_jwt,
        index_name=index_name,
        namespace=namespace,
        comment=upload_comment,
        batch_number=batch_number,
        db=db,
        request=request,
        topic_name=TXT_INGEST_TOPIC,
        log_prefix="[UPLOAD DOCUMENT]",
    )
