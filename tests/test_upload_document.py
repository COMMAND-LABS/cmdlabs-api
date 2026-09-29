"""
POST /api/vector-stores/upload-document: PDFs and Word documents are turned
into text and queued on the TXT ingest topic like a plain text upload.

GCS / Pub/Sub are external, so the upload service is replaced with a recorder;
knowledge-base authorization is covered by its own tests and stubbed here.
"""
import io

import docx
import fitz
import pytest

import src.routers.vectorStores.upload as upload_mod
from src.services.vector_stores_upload_service import TXT_INGEST_TOPIC, VectorStoresUploadService

URL = "/api/vector-stores/upload-document"
FORM = {"index_name": "kb-index", "namespace": "policies"}


@pytest.fixture()
def uploads(monkeypatch, test_account):
    """Capture what would have been sent to GCS / Pub/Sub."""
    calls = []

    async def fake_upload(self, *, file, topic_name, comment=None, **kwargs):
        body = await file.read()
        calls.append({
            "filename": file.filename,
            "content_type": file.content_type,
            "text": body.decode("utf-8"),
            "topic": topic_name,
            "comment": comment,
        })
        return {"success": True, "filename": file.filename, "gcs_bucket": "b", "gcs_file_path": "p"}

    logs = []

    def fake_log(db, **kwargs):
        logs.append(kwargs)
        return 1

    monkeypatch.setattr(VectorStoresUploadService, "upload_file_and_publish", fake_upload)
    monkeypatch.setattr(upload_mod, "record_ingestion_log", fake_log)
    monkeypatch.setattr(upload_mod, "_authorize_upload", lambda *a, **k: test_account.id)
    return {"calls": calls, "logs": logs}


def _pdf_bytes(text: str | None) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Refunds are processed within five days.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Plan"
    table.rows[0].cells[1].text = "Premium"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


async def test_pdf_text_is_extracted_and_uploaded(authed_client, uploads):
    files = {"file": ("Handbook.pdf", _pdf_bytes("Office hours are 9 to 5."), "application/pdf")}
    resp = await authed_client.post(URL, data=FORM, files=files)
    assert resp.status_code == 200, resp.text

    [call] = uploads["calls"]
    assert call["filename"] == "Handbook.txt"
    assert call["topic"] == TXT_INGEST_TOPIC
    assert "Office hours are 9 to 5." in call["text"]
    assert "Handbook.pdf" in call["comment"]
    assert uploads["logs"][0]["filenames"] == ["Handbook.txt"]


async def test_docx_paragraphs_and_tables_are_extracted(authed_client, uploads):
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    files = {"file": ("Policy.docx", _docx_bytes(), mime)}
    resp = await authed_client.post(URL, data={**FORM, "comment": "Q3"}, files=files)
    assert resp.status_code == 200, resp.text

    [call] = uploads["calls"]
    assert call["filename"] == "Policy.txt"
    assert call["topic"] == TXT_INGEST_TOPIC
    assert "Refunds are processed within five days." in call["text"]
    assert "Plan | Premium" in call["text"]
    assert call["comment"].startswith("Q3") and "Policy.docx" in call["comment"]


async def test_text_file_passes_through_unchanged(authed_client, uploads):
    files = {"file": ("notes.md", b"# Notes\nHello", "text/markdown")}
    resp = await authed_client.post(URL, data=FORM, files=files)
    assert resp.status_code == 200, resp.text

    [call] = uploads["calls"]
    assert call["filename"] == "notes.md"
    assert call["text"] == "# Notes\nHello"
    assert call["topic"] == TXT_INGEST_TOPIC


async def test_unsupported_type_is_rejected(authed_client, uploads):
    files = {"file": ("sheet.xlsx", b"whatever", "application/octet-stream")}
    resp = await authed_client.post(URL, data=FORM, files=files)
    assert resp.status_code == 400
    assert uploads["calls"] == []


async def test_pdf_without_text_is_rejected(authed_client, uploads):
    files = {"file": ("scan.pdf", _pdf_bytes(None), "application/pdf")}
    resp = await authed_client.post(URL, data=FORM, files=files)
    assert resp.status_code == 400
    assert "couldn't find any text" in resp.json()["detail"]
    assert uploads["calls"] == []


async def test_corrupt_pdf_gives_fixed_message(authed_client, uploads):
    files = {"file": ("broken.pdf", b"not a pdf at all", "application/pdf")}
    resp = await authed_client.post(URL, data=FORM, files=files)
    assert resp.status_code == 400
    assert "couldn't read this file" in resp.json()["detail"]
    assert uploads["calls"] == []


@pytest.mark.parametrize("route, filename", [
    ("/api/vector-stores/upload-document", "notes.txt"),
    ("/api/vector-stores/upload-csv", "faq.csv"),
])
async def test_empty_namespace_targets_main_section(authed_client, uploads, route, filename):
    """The default ("Main") section is namespace "" — a valid upload target."""
    files = {"file": (filename, b"q,a\nhi,there", "text/plain")}
    resp = await authed_client.post(route, data={"index_name": "kb-index", "namespace": ""}, files=files)
    assert resp.status_code == 200, resp.text
    assert len(uploads["calls"]) == 1
